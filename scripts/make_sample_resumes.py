"""Generate 3 SYNTHETIC test resumes in samples/resumes/ (fictional people, example.com emails).

They are placeholders until the real test resumes are provided, and are designed
to land in different Stage 2 outcomes against config/job_description.yaml:
  1. strong match  2. borderline/partial  3. weak match (different stack)
Layouts deliberately use a page header and tables, like real resume templates.
Run:  uv run python scripts/make_sample_resumes.py
"""

from pathlib import Path

from docx import Document
from docx.shared import Pt

OUT = Path(__file__).resolve().parents[1] / "samples" / "resumes"


def _doc(name: str, contact: str, headline: str):
    d = Document()
    d.styles["Normal"].font.size = Pt(10.5)
    hdr = d.sections[0].header.paragraphs[0]
    hdr.text = contact  # contact details in the page header, as many templates do
    d.add_heading(name, level=0)
    d.add_paragraph(headline).runs[0].italic = True
    return d


def _skills_table(d, rows):
    d.add_heading("Skills", level=1)
    t = d.add_table(rows=0, cols=2)
    t.style = "Table Grid"
    for k, v in rows:
        c = t.add_row().cells
        c[0].text, c[1].text = k, v


def _job(d, title, when, bullets):
    d.add_paragraph().add_run(title).bold = True
    d.add_paragraph(when)
    for b in bullets:
        d.add_paragraph(b, style="List Bullet")


def strong():
    d = _doc("Aarav Sharma", "aarav.sharma@example.com | +91 90000 00001 | Hyderabad, India | linkedin.com/in/aarav-sharma-example",
             "Senior Python Backend Engineer — 6 years building REST APIs")
    d.add_heading("Summary", level=1)
    d.add_paragraph("Backend engineer with 6 years of experience designing and operating Python services and REST APIs "
                    "on AWS. Comfortable owning features end to end, from schema design to CI/CD and monitoring.")
    _skills_table(d, [("Languages", "Python, SQL, Bash"),
                      ("Frameworks", "FastAPI, Django, Django REST Framework, Celery"),
                      ("Data", "PostgreSQL, Redis, SQLAlchemy"),
                      ("DevOps", "Docker, GitHub Actions (CI/CD), AWS (ECS, Lambda, RDS, S3), Terraform"),
                      ("Practices", "Git, code review, pytest, TDD, agile/Scrum")])
    d.add_heading("Experience", level=1)
    _job(d, "Senior Software Engineer — Finlytics Pvt Ltd, Hyderabad", "Apr 2022 – Present",
         ["Lead developer of a FastAPI payments API serving 2M requests/day on AWS ECS.",
          "Redesigned PostgreSQL schema and queries, cutting p95 latency from 480ms to 120ms.",
          "Introduced GitHub Actions CI/CD with automated pytest suites (85% coverage)."])
    _job(d, "Software Engineer — CloudKart Technologies, Bengaluru", "Jul 2019 – Mar 2022",
         ["Built Django REST Framework services for order management.",
          "Containerised 12 services with Docker; migrated deployments to AWS.",
          "Mentored 3 junior engineers; ran weekly code reviews."])
    d.add_heading("Education", level=1)
    d.add_paragraph("B.Tech, Computer Science and Engineering — JNTU Hyderabad, 2019")
    d.add_heading("Certifications", level=1)
    d.add_paragraph("AWS Certified Developer – Associate (2023)")
    d.add_heading("Languages", level=1)
    d.add_paragraph("English, Hindi, Telugu")
    return d


def partial():
    d = _doc("Priya Nair", "priya.nair@example.com | +91 90000 00002 | Kochi, Kerala",
             "Data Analyst moving into backend development")
    d.add_heading("Profile", level=1)
    d.add_paragraph("Data analyst with 2.5 years of experience automating reporting with Python and SQL. "
                    "Built a few internal Flask endpoints; learning FastAPI and Docker.")
    _skills_table(d, [("Programming", "Python (pandas, NumPy), SQL"),
                      ("Tools", "Power BI, Excel, Jupyter, Git"),
                      ("Databases", "MySQL, PostgreSQL (basic)"),
                      ("Learning", "FastAPI, Docker")])
    d.add_heading("Work Experience", level=1)
    _job(d, "Data Analyst — Coastal Retail Analytics, Kochi", "Jan 2023 – Present",
         ["Automated 20+ weekly Excel reports using Python and SQL, saving ~15 hours/week.",
          "Wrote two small Flask endpoints exposing sales metrics to the BI team.",
          "Maintained scripts in a shared Git repository."])
    _job(d, "Analyst Intern — Coastal Retail Analytics, Kochi", "Jun 2022 – Dec 2022",
         ["Cleaned and validated sales data; built Power BI dashboards."])
    d.add_heading("Education", level=1)
    d.add_paragraph("B.Sc. Mathematics — University of Kerala, 2022")
    return d


def weak():
    d = _doc("Rohan Mehta", "rohan.mehta@example.com | +91 90000 00003 | Pune, Maharashtra",
             "Frontend Developer — React & TypeScript")
    d.add_heading("About Me", level=1)
    d.add_paragraph("Frontend developer with 4 years of experience building responsive web apps in React and TypeScript. "
                    "I consume REST APIs daily and enjoy UI performance work.")
    _skills_table(d, [("Frontend", "React, TypeScript, JavaScript, Redux, HTML5, CSS3, Tailwind"),
                      ("Testing", "Jest, React Testing Library, Cypress"),
                      ("Tools", "Git, Webpack, Vite, Figma")])
    d.add_heading("Experience", level=1)
    _job(d, "Frontend Developer — PixelCraft Studios, Pune", "Aug 2021 – Present",
         ["Built a React/TypeScript design system used across 5 products.",
          "Integrated REST APIs from backend teams; improved Lighthouse score from 62 to 94."])
    _job(d, "Junior Web Developer — WebNest Solutions, Pune", "Jul 2020 – Jul 2021",
         ["Developed marketing sites with HTML, CSS and vanilla JavaScript."])
    d.add_heading("Education", level=1)
    d.add_paragraph("B.E. Information Technology — Savitribai Phule Pune University, 2020")
    return d


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    for fname, builder in (("aarav_sharma.docx", strong), ("priya_nair.docx", partial), ("rohan_mehta.docx", weak)):
        builder().save(OUT / fname)
        print(f"wrote samples/resumes/{fname}")


if __name__ == "__main__":
    main()
