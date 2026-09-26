"""Unit tests: docx reader, config validation, prompts, LLM parsing, Gemini schema, scoring."""

from __future__ import annotations

import copy
import json

import pytest
from conftest import ASSESSMENTS, SAMPLES
from docx import Document
from lxml import etree

from screening.config import CRITERIA, ConfigError, Scoring, load_config
from screening.docx_reader import DocxReadError, read_docx
from screening.llm.base import LLMResponseError, parse_json_response
from screening.llm.gemini import to_gemini_schema
from screening.paths import PROMPTS_DIR
from screening.prompts import PromptError, load_prompt
from screening.schemas import ResumeExtractionLLM, ShortlistAssessmentLLM, TranscriptParseLLM
from screening.stage2_shortlist import blind_profile, decide, reconcile_skills

# ---------------------------------------------------------------- docx reader


def test_reader_gets_header_table_and_unicode():
    text = read_docx(SAMPLES / "resumes" / "aarav_sharma.docx")
    lines = text.splitlines()
    assert lines[0].startswith("aarav.sharma@example.com | +91 90000 00001")  # page header first
    assert "Frameworks | FastAPI, Django, Django REST Framework, Celery" in lines  # table row
    assert "—" in text  # em dash preserved


def _textbox_xml(text):
    w = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    box = f'<w:txbxContent xmlns:w="{w}"><w:p><w:r><w:t>{text}</w:t></w:r></w:p></w:txbxContent>'
    return (
        f'<w:r xmlns:w="{w}" xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
        'xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape" '
        'xmlns:v="urn:schemas-microsoft-com:vml">'
        f'<mc:AlternateContent><mc:Choice Requires="wps"><w:drawing><wps:txbx>{box}</wps:txbx></w:drawing></mc:Choice>'
        f'<mc:Fallback><w:pict><v:textbox>{box}</v:textbox></w:pict></mc:Fallback></mc:AlternateContent></w:r>'
    )


def test_reader_textbox_once_and_merged_cells(tmp_path):
    d = Document()
    p = d.add_paragraph("Body")
    p._p.append(etree.fromstring(_textbox_xml("Phone: +91 12345")))
    t = d.add_table(rows=1, cols=3)
    merged = t.cell(0, 0).merge(t.cell(0, 1))
    merged.text = "Merged"
    t.cell(0, 2).text = "Right"
    f = tmp_path / "x.docx"
    d.save(f)
    text = read_docx(f)
    assert text.count("Phone: +91 12345") == 1  # DrawingML + VML fallback -> one copy
    assert "Merged | Right" in text.splitlines()


def test_reader_rejects_bad_files(tmp_path):
    (tmp_path / "a.docx").write_text("nope")
    (tmp_path / "b.pdf").write_text("nope")
    with pytest.raises(DocxReadError, match="Not a readable"):
        read_docx(tmp_path / "a.docx")
    with pytest.raises(DocxReadError, match="Unsupported"):
        read_docx(tmp_path / "b.pdf")


# ---------------------------------------------------------------- config


def test_real_config_loads():
    cfg = load_config()
    assert cfg.job.must_have_skills and cfg.questions.questions
    assert set(cfg.settings.scoring.weights) == set(CRITERIA)


@pytest.mark.parametrize("weights", [
    {"must_have_skills": 0.5, "experience": 0.3, "nice_to_have_skills": 0.15, "role_relevance": 0.15},
    {"must_have_skills": 0.7, "experience": 0.3},
])
def test_bad_weights_rejected(weights):
    with pytest.raises(ValueError):
        Scoring(weights=weights)


def test_missing_api_key(monkeypatch):
    cfg = load_config()
    monkeypatch.delenv(cfg.settings.llm.api_key_env, raising=False)
    with pytest.raises(ConfigError, match="not set"):
        cfg.api_key()


# ---------------------------------------------------------------- prompts


def test_every_prompt_file_loads():
    files = sorted(PROMPTS_DIR.glob("*/*.md"))
    assert len(files) == 14
    for f in files:
        assert load_prompt(f.parent.name, f.name).text.strip()


def test_prompt_render_strict_and_single_pass():
    p = load_prompt("stage1_extraction", "extract_resume.md")
    with pytest.raises(PromptError):
        p.render(resume_text="x")
    out = p.render(resume_text="I wrote {{today}} in my CV", today="2026-01-01")
    assert "I wrote {{today}} in my CV" in out  # user content is never expanded


# ---------------------------------------------------------------- LLM parsing / schema


def test_parse_json_response():
    good = json.dumps(ASSESSMENTS["Data Analyst"])
    assert parse_json_response(f"```json\n{good}\n```", ShortlistAssessmentLLM).rationale == "Partial."
    with pytest.raises(LLMResponseError, match="not valid JSON"):
        parse_json_response("{oops", ShortlistAssessmentLLM)
    bad = copy.deepcopy(ASSESSMENTS["Data Analyst"])
    bad["experience"]["score"] = 150
    with pytest.raises(LLMResponseError, match="does not match"):
        parse_json_response(json.dumps(bad), ShortlistAssessmentLLM)
    with pytest.raises(LLMResponseError, match="empty"):
        parse_json_response("  ", ShortlistAssessmentLLM)


@pytest.mark.parametrize("model", [ResumeExtractionLLM, ShortlistAssessmentLLM, TranscriptParseLLM])
def test_gemini_schema_is_plain_openapi(model):
    def walk(node):
        yield node
        for k, v in node.items():
            children = v.values() if k == "properties" else [v] if isinstance(v, dict) else []
            for c in children:
                yield from walk(c)

    nodes = list(walk(to_gemini_schema(model)))
    for n in nodes:  # schema keywords only; property *names* like "title" are fine
        assert not {"$ref", "$defs", "anyOf", "additionalProperties", "title", "default"} & n.keys(), n
    nullable_fields = {k for k, f in model.model_fields.items() if not f.is_required() or "None" in str(f.annotation)}
    root_props = nodes[0]["properties"]
    assert all(root_props[k].get("nullable") for k in nullable_fields)


# ---------------------------------------------------------------- scoring


def test_decide_threshold_and_must_have_rule():
    cfg = load_config()
    a = ShortlistAssessmentLLM.model_validate(ASSESSMENTS["Senior Software Engineer"])
    criteria, overall, decision, _ = decide(a, cfg.settings, [])
    assert overall == 89.75 and decision == "shortlisted"
    assert [c.criterion for c in criteria] == list(CRITERIA)

    strict = cfg.settings.model_copy(deep=True)
    strict.thresholds.require_all_must_have = True
    _, _, decision, reasons = decide(a, strict, ["SQL"])
    assert decision == "rejected" and "SQL" in reasons[-1]

    edge = cfg.settings.model_copy(deep=True)
    edge.thresholds.shortlist_score = 89.75
    assert decide(a, edge, [])[2] == "shortlisted"  # >= is inclusive


def test_reconcile_skills():
    warnings = []
    m, miss = reconcile_skills(["python", "rest-api development", "Rust"], ["SQL", "python"],
                               ["Python", "REST API development", "SQL", "Git"], warnings.append)
    assert m == ["REST API development"]
    assert miss == ["Python", "SQL", "Git"]
    assert len(warnings) == 2  # unknown 'rust' + python in both lists


def test_blind_profile_hides_identity():
    from conftest import PROFILES

    out = blind_profile(ResumeExtractionLLM.model_validate(PROFILES["Aarav Sharma"]))
    for secret in ("Aarav", "+91", "@example.com", "Hyderabad"):
        assert secret not in out
    assert "FastAPI" in out
